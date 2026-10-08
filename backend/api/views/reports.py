"""Reports (sales, purchases, stock, outstanding, valuation, trends ...) and Excel/PDF exports."""
import io
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal

from django.conf import settings as dj
from django.db.models import Sum
from django.http import HttpResponse
from rest_framework.decorators import api_view
from rest_framework.response import Response

from ..core.fieldtypes import utcnow
from ..models import (Categories, CustomerPayments, Customers, GoodsReceiptItems, GoodsReceipts, InvoiceItems, Invoices, Products, PurchaseOrderItems, PurchaseOrders,
                      PurchaseReturns, SalesReturns, Stock, StockMovements, Suppliers, SupplierPayments, Warehouses)


def F(v): return float(v or 0)
def iso(v): return v.isoformat() if v else None
def names(model, ids, attr="name"):
    ids = {i for i in ids if i}
    return {o.pk: getattr(o, attr, None) for o in model.objects.filter(pk__in=ids)} if ids else {}


def q_dates(request):
    qp = request.query_params
    def p(k, end=False):
        v = qp.get(k)
        if not v: return None
        try: d = datetime.fromisoformat(v[:10]); return d + timedelta(days=1) - timedelta(seconds=1) if end else d
        except ValueError: return None
    return p("fromDate") or p("startDate") or p("from"), p("toDate", True) or p("endDate", True) or p("to", True)


def qint(request, key):
    v = request.query_params.get(key)
    return int(v) if v and v.isdigit() else None


def sales_rows(request, include_cancelled=False):
    frm, to = q_dates(request)
    q = Invoices.objects.filter(is_deleted=False)
    if not include_cancelled: q = q.filter(is_cancelled=False)
    if frm: q = q.filter(invoice_date__gte=frm)
    if to: q = q.filter(invoice_date__lte=to)
    for k, a in (("warehouseId", "warehouse_id"), ("customerId", "customer_id")):
        if qint(request, k): q = q.filter(**{a: qint(request, k)})
    if request.query_params.get("status") and request.query_params["status"].lower() != "all": q = q.filter(status__iexact=request.query_params["status"])
    invs = list(q.order_by("-invoice_date", "-invoice_id"))
    items = defaultdict(list)
    for it in InvoiceItems.objects.filter(invoice_id__in=[i.pk for i in invs]): items[it.invoice_id].append(it)
    pid = qint(request, "productId"); cat = qint(request, "categoryId")
    pm = {p.pk: p for p in Products.objects.filter(pk__in={x.product_id for l in items.values() for x in l})}
    cn = names(Customers, [i.customer_id for i in invs]); wn = names(Warehouses, [i.warehouse_id for i in invs]); out = []
    for i in invs:
        its = [x for x in items[i.pk] if (not pid or x.product_id == pid) and (not cat or (pm.get(x.product_id) and pm[x.product_id].category_id == cat))]
        if (pid or cat) and not its: continue
        out.append({"id": i.pk, "soId": i.so_id or i.pk, "invoiceId": i.pk, "customerId": i.customer_id, "customer": cn.get(i.customer_id), "customerName": cn.get(i.customer_id), "soNumber": i.invoice_number,
                    "invoiceNumber": i.invoice_number, "orderDate": iso(i.invoice_date), "invoiceDate": iso(i.invoice_date), "totalAmount": F(i.total_amount), "paidAmount": F(i.paid_amount),
                    "balanceAmount": F(i.balance_amount), "status": i.status, "warehouseId": i.warehouse_id, "warehouseName": wn.get(i.warehouse_id), "warehouse": wn.get(i.warehouse_id),
                    "items": [{"productId": x.product_id, "productName": pm[x.product_id].name if x.product_id in pm else None, "product": pm[x.product_id].name if x.product_id in pm else None,
                               "sku": pm[x.product_id].sku if x.product_id in pm else None, "quantity": F(x.quantity), "price": F(x.price), "total": F(x.total), "taxAmount": F(x.tax_amount),
                               "costPrice": F(pm[x.product_id].cost_price) if x.product_id in pm else 0} for x in its]})
    return out


@api_view(["GET"])
def sales(request): return Response(sales_rows(request))


@api_view(["GET"])
def invoices(request):
    return Response([{k: v for k, v in r.items() if k not in ("soId", "soNumber", "orderDate")} for r in sales_rows(request, include_cancelled=True)])


@api_view(["GET"])
def purchases(request):
    return Response(purchase_rows(request))


def purchase_rows(request):
    frm, to = q_dates(request); q = PurchaseOrders.objects.filter(is_deleted=False, is_cancelled=False)
    if frm: q = q.filter(order_date__gte=frm)
    if to: q = q.filter(order_date__lte=to)
    if qint(request, "supplierId"): q = q.filter(supplier_id=qint(request, "supplierId"))
    if request.query_params.get("status") and request.query_params["status"].lower() != "all": q = q.filter(status__iexact=request.query_params["status"])
    pos = list(q.order_by("-order_date", "-po_id")); items = defaultdict(list)
    for it in PurchaseOrderItems.objects.filter(po_id__in=[p.pk for p in pos]): items[it.po_id].append(it)
    wh = {}
    for g in GoodsReceipts.objects.filter(po_id__in=[p.pk for p in pos], is_cancelled=False): wh.setdefault(g.po_id, g.warehouse_id)
    wid = qint(request, "warehouseId"); pid = qint(request, "productId")
    sn = names(Suppliers, [p.supplier_id for p in pos]); wn = names(Warehouses, wh.values()); pn = names(Products, [x.product_id for l in items.values() for x in l]); out = []
    for p in pos:
        if wid and wh.get(p.pk) != wid: continue
        its = [x for x in items[p.pk] if not pid or x.product_id == pid]
        if pid and not its: continue
        w = wh.get(p.pk)
        out.append({"id": p.pk, "poId": p.pk, "poNumber": p.po_number, "supplierId": p.supplier_id, "supplier": sn.get(p.supplier_id), "supplierName": sn.get(p.supplier_id), "orderDate": iso(p.order_date),
                    "totalAmount": F(p.total_amount), "status": p.status, "warehouseId": w, "warehouseName": wn.get(w), "warehouse": wn.get(w),
                    "items": [{"productId": x.product_id, "productName": pn.get(x.product_id), "product": pn.get(x.product_id), "quantity": F(x.quantity), "receivedQuantity": F(x.received_quantity), "price": F(x.price), "total": F(x.total),
                               "warehouseId": w, "warehouseName": wn.get(w), "warehouse": wn.get(w)} for x in its]})
    return out


def stock_rows(request):
    q = Stock.objects.filter(is_deleted=False)
    if qint(request, "warehouseId"): q = q.filter(warehouse_id=qint(request, "warehouseId"))
    if qint(request, "productId"): q = q.filter(product_id=qint(request, "productId"))
    rows = list(q); pm = {p.pk: p for p in Products.objects.filter(pk__in={r.product_id for r in rows}, is_deleted=False)}
    cat = qint(request, "categoryId"); cn = names(Categories, [p.category_id for p in pm.values()]); wn = names(Warehouses, [r.warehouse_id for r in rows]); out = []
    for r in rows:
        p = pm.get(r.product_id)
        if not p or (cat and p.category_id != cat): continue
        lvl = p.reorder_level or dj.LOW_STOCK_THRESHOLD; av = F(r.available_quantity if r.available_quantity is not None else r.quantity)
        out.append({"id": r.pk, "stockId": r.pk, "productId": p.pk, "product": p.name, "productName": p.name, "name": p.name, "sku": p.sku, "categoryId": p.category_id, "category": cn.get(p.category_id),
                    "categoryName": cn.get(p.category_id), "warehouseId": r.warehouse_id, "warehouseName": wn.get(r.warehouse_id), "warehouse": wn.get(r.warehouse_id), "price": F(p.price), "costPrice": F(p.cost_price),
                    "quantity": F(r.quantity), "reservedQuantity": F(r.reserved_quantity), "availableQuantity": av, "status": "Out of Stock" if av <= 0 else ("Low Stock" if av <= lvl else "In Stock"), "reorderLevel": lvl})
    return out


@api_view(["GET"])
def stock(request): return Response(stock_rows(request))


@api_view(["GET"])
def customer_balances(request):
    q = Customers.objects.all()
    if qint(request, "customerId"): q = q.filter(pk=qint(request, "customerId"))
    return Response([{"id": c.pk, "customerId": c.pk, "name": c.name, "customer": c.name, "customerName": c.name, "company": c.company, "creditLimit": F(c.credit_limit), "outstandingBalance": F(c.outstanding_balance),
                      "status": c.status, "warehouseId": None, "warehouseName": None, "warehouse": None} for c in q.order_by("-outstanding_balance")])


# ---------------------------------------------------------------- filters
@api_view(["GET"])
def f_warehouses(request): return Response([{"id": w.pk, "warehouseId": w.pk, "name": w.name, "warehouseName": w.name} for w in Warehouses.objects.filter(is_deleted=False).order_by("name")])
@api_view(["GET"])
def f_categories(request): return Response([{"id": c.pk, "categoryId": c.pk, "name": c.name, "categoryName": c.name} for c in Categories.objects.filter(is_deleted=False).order_by("name")])
@api_view(["GET"])
def f_products(request): return Response([{"id": p.pk, "productId": p.pk, "name": p.name, "productName": p.name, "sku": p.sku, "categoryId": p.category_id} for p in Products.objects.filter(is_deleted=False).order_by("name")])
@api_view(["GET"])
def f_customers(request): return Response([{"id": c.pk, "customerId": c.pk, "name": c.name, "customerName": c.name} for c in Customers.objects.order_by("name")])
@api_view(["GET"])
def f_suppliers(request): return Response([{"id": s.pk, "supplierId": s.pk, "name": s.name, "supplierName": s.name} for s in Suppliers.objects.filter(is_deleted=False).order_by("name")])


# ---------------------------------------------------------------- analytics
@api_view(["GET"])
def transaction_trend(request):
    months = []; d = utcnow().replace(day=1, hour=0, minute=0, second=0)
    for _ in range(6):
        months.append(d); d = (d - timedelta(days=1)).replace(day=1)
    months.reverse(); out = []
    for m in months:
        nxt = (m + timedelta(days=32)).replace(day=1)
        s = Invoices.objects.filter(is_deleted=False, is_cancelled=False, invoice_date__gte=m, invoice_date__lt=nxt).aggregate(t=Sum("total_amount"))["t"]
        p = PurchaseOrders.objects.filter(is_deleted=False, is_cancelled=False, order_date__gte=m, order_date__lt=nxt).aggregate(t=Sum("total_amount"))["t"]
        out.append({"month": m.strftime("%Y-%m"), "monthName": m.strftime("%b"), "sales": F(s), "purchases": F(p)})
    return Response(out)


@api_view(["GET"])
def stock_availability(request):
    rows = stock_rows(request)
    return Response({"inStock": sum(1 for r in rows if r["status"] == "In Stock"), "lowStock": sum(1 for r in rows if r["status"] == "Low Stock"), "outOfStock": sum(1 for r in rows if r["status"] == "Out of Stock")})


@api_view(["GET"])
def top_customers(request):
    g = {}
    for r in sales_rows(request):
        if not r["customerId"]: continue
        x = g.setdefault(r["customerId"], {"customerId": r["customerId"], "customerName": r["customerName"], "totalOrders": 0, "totalSalesValue": 0.0, "outstandingAmount": 0.0, "lastPurchaseDate": None,
                                          "warehouseId": r["warehouseId"], "warehouseName": r["warehouseName"], "warehouse": r["warehouseName"]})
        x["totalOrders"] += 1; x["totalSalesValue"] += r["totalAmount"]; x["outstandingAmount"] += r["balanceAmount"]; x["lastPurchaseDate"] = max(filter(None, [x["lastPurchaseDate"], r["invoiceDate"]]), default=None)
    return Response(sorted(g.values(), key=lambda x: -x["totalSalesValue"]))


@api_view(["GET"])
def top_suppliers(request):
    g = {}
    for r in purchase_rows(request):
        x = g.setdefault(r["supplierId"], {"supplierId": r["supplierId"], "supplierName": r["supplierName"], "totalOrders": 0, "totalPurchaseValue": 0.0, "lastOrderDate": None})
        x["totalOrders"] += 1; x["totalPurchaseValue"] += r["totalAmount"]; x["lastOrderDate"] = max(filter(None, [x["lastOrderDate"], r["orderDate"]]), default=None)
    return Response(sorted(g.values(), key=lambda x: -x["totalPurchaseValue"]))


@api_view(["GET"])
def customer_outstanding(request):
    out = []; today = utcnow().date()
    for r in sales_rows(request):
        if r["balanceAmount"] <= 0: continue
        due = Invoices.objects.filter(pk=r["invoiceId"]).values_list("due_date", flat=True).first(); dd = due.date() if isinstance(due, datetime) else due
        late = (today - dd).days if dd else 0
        out.append({"id": r["invoiceId"], "customerId": r["customerId"], "customerName": r["customerName"], "invoiceNumber": r["invoiceNumber"], "invoiceDate": r["invoiceDate"], "dueDate": iso(dd), "invoiceAmount": r["totalAmount"],
                    "paidAmount": r["paidAmount"], "balanceAmount": r["balanceAmount"], "warehouseId": r["warehouseId"], "warehouseName": r["warehouseName"], "warehouse": r["warehouseName"],
                    "agingStatus": "Current" if late <= 0 else ("1-30 days" if late <= 30 else ("31-60 days" if late <= 60 else "60+ days"))})
    return Response(sorted(out, key=lambda x: -x["balanceAmount"]))


@api_view(["GET"])
def supplier_outstanding(request):
    return Response(_supplier_outstanding(request))


def _supplier_outstanding(request):
    out = []
    for r in purchase_rows(request):
        paid = F(SupplierPayments.objects.filter(po_id=r["poId"], is_cancelled=False).aggregate(t=Sum("amount"))["t"])
        bal = max(r["totalAmount"] - paid, 0)
        if bal > 0: out.append({"id": r["poId"], "supplierId": r["supplierId"], "supplierName": r["supplierName"], "poNumber": r["poNumber"], "orderDate": r["orderDate"], "totalAmount": r["totalAmount"], "paid": paid, "paidAmount": paid, "balanceAmount": bal})
    return sorted(out, key=lambda x: -x["balanceAmount"])


@api_view(["GET"])
def inventory_valuation(request):
    rows = stock_rows(request); last = {}
    for g in GoodsReceipts.objects.filter(is_cancelled=False).order_by("receipt_date"):
        for it in GoodsReceiptItems.objects.filter(grn_id=g.pk): last[it.product_id] = g.receipt_date
    return Response([{"id": r["stockId"], "productId": r["productId"], "productName": r["productName"], "sku": r["sku"], "categoryId": r["categoryId"], "category": r["category"], "warehouseId": r["warehouseId"],
                      "warehouseName": r["warehouseName"], "warehouse": r["warehouseName"], "quantityAvailable": r["availableQuantity"], "averageCost": r["costPrice"], "totalStockValue": round(r["availableQuantity"] * r["costPrice"], 2),
                      "lastPurchaseDate": iso(last.get(r["productId"]))} for r in rows])


@api_view(["GET"])
def low_stock(request):
    return Response(_low_stock(request))


def _low_stock(request):
    out = []
    for r in stock_rows(request):
        if r["availableQuantity"] <= r["reorderLevel"]:
            out.append({"id": r["stockId"], "productId": r["productId"], "productName": r["productName"], "sku": r["sku"], "categoryId": r["categoryId"], "category": r["category"], "availableStock": r["availableQuantity"],
                        "minimumStockLevel": r["reorderLevel"], "reorderQuantity": max(r["reorderLevel"] * 2 - r["availableQuantity"], 0), "warehouseId": r["warehouseId"], "warehouseName": r["warehouseName"], "warehouse": r["warehouseName"]})
    return sorted(out, key=lambda x: x["availableStock"])


def _sold_by_product(days=None):
    q = InvoiceItems.objects.filter(invoice_id__in=Invoices.objects.filter(is_deleted=False, is_cancelled=False).values("invoice_id"))
    sold = defaultdict(lambda: [0.0, 0.0, None]); inv = {i.pk: i.invoice_date for i in Invoices.objects.filter(is_deleted=False, is_cancelled=False)}
    for it in q:
        d = inv.get(it.invoice_id)
        if days and d and d < utcnow() - timedelta(days=days): continue
        s = sold[it.product_id]; s[0] += F(it.quantity); s[1] += F(it.total); s[2] = max(filter(None, [s[2], d]), default=None)
    return sold


@api_view(["GET"])
def fast_moving(request):
    return Response(_fast_moving(request))


def _fast_moving(request):
    sold = _sold_by_product(90); pm = {p.pk: p for p in Products.objects.filter(pk__in=sold.keys())}
    return sorted([{"productId": k, "productName": pm[k].name, "sku": pm[k].sku, "quantitySold": v[0], "salesValue": v[1], "lastSaleDate": iso(v[2])} for k, v in sold.items() if k in pm], key=lambda x: -x["quantitySold"])[:50]


@api_view(["GET"])
def slow_moving(request):
    sold = _sold_by_product(); out = []; now = utcnow()
    for r in stock_rows(request):
        if r["quantity"] <= 0: continue
        last = sold[r["productId"]][2] if r["productId"] in sold else None
        out.append({"productId": r["productId"], "productName": r["productName"], "sku": r["sku"], "quantityInStock": r["quantity"], "lastSaleDate": iso(last), "daysSinceLastSale": (now - last).days if last else 9999,
                    "warehouseId": r["warehouseId"], "warehouseName": r["warehouseName"]})
    return Response(sorted(out, key=lambda x: -x["daysSinceLastSale"]))


@api_view(["GET"])
def profitability(request):
    g = {}
    for r in sales_rows(request):
        for it in r["items"]:
            x = g.setdefault(it["productId"], {"productId": it["productId"], "productName": it["productName"], "sku": it["sku"], "quantitySold": 0.0, "revenue": 0.0, "cost": 0.0})
            x["quantitySold"] += it["quantity"]; x["revenue"] += it["total"] - it["taxAmount"]; x["cost"] += it["quantity"] * it["costPrice"]
    for x in g.values(): x["profit"] = round(x["revenue"] - x["cost"], 2); x["margin"] = round(x["profit"] / x["revenue"] * 100, 2) if x["revenue"] else 0
    return Response(sorted(g.values(), key=lambda x: -x["profit"]))


@api_view(["GET"])
def gst_tax(request):
    m = defaultdict(lambda: [0.0, 0.0])
    for r in sales_rows(request):
        k = (r["invoiceDate"] or "")[:7]; m[k][0] += r["totalAmount"]; m[k][1] += sum(i["taxAmount"] for i in r["items"])
    return Response([{"date": k, "total": round(v[0], 2), "tax": round(v[1], 2)} for k, v in sorted(m.items()) if k])


@api_view(["GET"])
def forecasting(request):
    sold = _sold_by_product(60); pm = {p.pk: p for p in Products.objects.filter(pk__in=sold.keys())}; out = []; n = 1
    for k, v in sorted(sold.items(), key=lambda kv: -kv[1][0])[:5]:
        if k in pm:
            per_month = v[0] / 2
            out.append({"id": n, "insight": f"{pm[k].name} sells about {per_month:.0f} units per month.", "prediction": f"Expect roughly {per_month * 1.05:.0f} units next month."}); n += 1
    return Response(out or [{"id": 1, "insight": "Not enough sales history yet.", "prediction": "Create invoices to see demand forecasts."}])


@api_view(["GET"])
def warehouse_performance(request):
    out = []
    for w in Warehouses.objects.filter(is_deleted=False):
        rows = [r for r in stock_rows(request) if r["warehouseId"] == w.pk]
        out.append({"id": w.pk, "warehouseId": w.pk, "warehouseName": w.name, "stockValue": round(sum(r["quantity"] * r["costPrice"] for r in rows), 2), "totalProducts": len(rows),
                    "lowStockItems": sum(1 for r in rows if r["status"] != "In Stock"), "damagedItems": 0,
                    "salesDispatches": Invoices.objects.filter(warehouse_id=w.pk, is_deleted=False, is_cancelled=False).count(), "purchaseReceipts": GoodsReceipts.objects.filter(warehouse_id=w.pk, is_cancelled=False, is_deleted=False).count()})
    return Response(out)


@api_view(["GET"])
def summary(request):
    s = sales_rows(request); p = purchase_rows(request); low = _low_stock(request); top = _fast_moving(request)
    ts, tp = sum(r["totalAmount"] for r in s), sum(r["totalAmount"] for r in p)
    pay = sum(r["balanceAmount"] for r in _supplier_outstanding(request))
    return Response({"totalSales": round(ts, 2), "totalPurchases": round(tp, 2), "lowStockItems": len(low), "payables": round(pay, 2), "profit": round(ts - tp, 2), "topSellingItem": top[0]["productName"] if top else "No sales yet"})


@api_view(["GET"])
def returns_report(request):
    out = [{"id": r.pk, "returnNumber": r.return_number, "type": "Sales", "date": iso(r.return_date), "status": r.status, "amount": F(r.grand_total), "reason": r.reason} for r in SalesReturns.objects.filter(is_deleted=False)]
    out += [{"id": r.pk, "returnNumber": r.return_number, "type": "Purchase", "date": iso(r.return_date), "status": r.status, "amount": F(r.total_return_amount), "reason": r.reason} for r in PurchaseReturns.objects.all()]
    return Response(sorted(out, key=lambda x: x["date"] or "", reverse=True))


@api_view(["GET"])
def empty_report(request): return Response([])


# ---------------------------------------------------------------- exports
def _xlsx(title, headers, rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = title[:30]; ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="5A49D6")
    for r in rows: ws.append(r)
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 3, 42)
    buf = io.BytesIO(); wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); resp["Content-Disposition"] = f'attachment; filename="{filename}"'; return resp


def _pdf(title, headers, rows, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO(); doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24); st = getSampleStyleSheet()
    t = Table([headers] + [[str(c)[:30] for c in r] for r in rows], repeatRows=1); t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5a49d6")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                                                                                       ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c9c5ee")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    doc.build([Paragraph(f"<b>{title}</b> &nbsp; <font size=8>Generated {utcnow():%Y-%m-%d %H:%M} UTC</font>", st["Title"]), Spacer(1, 8), t])
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf"); resp["Content-Disposition"] = f'attachment; filename="{filename}"'; return resp


SALES_HEAD = ["Invoice", "Date", "Customer", "Warehouse", "Status", "Total", "Paid", "Balance"]
STOCK_HEAD = ["Product", "SKU", "Category", "Warehouse", "Quantity", "Reserved", "Available", "Status"]


def _sales_table(request): return [[r["invoiceNumber"], (r["invoiceDate"] or "")[:10], r["customerName"], r["warehouseName"], r["status"], r["totalAmount"], r["paidAmount"], r["balanceAmount"]] for r in sales_rows(request)]
def _stock_table(request): return [[r["productName"], r["sku"], r["category"], r["warehouseName"], r["quantity"], r["reservedQuantity"], r["availableQuantity"], r["status"]] for r in stock_rows(request)]


@api_view(["GET"])
def export_sales(request): return _xlsx("Sales", SALES_HEAD, _sales_table(request), f"sales-report-{utcnow():%Y%m%d}.xlsx")
@api_view(["GET"])
def export_stock(request): return _xlsx("Stock", STOCK_HEAD, _stock_table(request), f"stock-report-{utcnow():%Y%m%d}.xlsx")
@api_view(["GET"])
def export_sales_pdf(request): return _pdf("Sales Report", SALES_HEAD, _sales_table(request), f"sales-report-{utcnow():%Y%m%d}.pdf")
@api_view(["GET"])
def export_stock_pdf(request): return _pdf("Stock Report", STOCK_HEAD, _stock_table(request), f"stock-report-{utcnow():%Y%m%d}.pdf")
